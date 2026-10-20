import argparse
import json
import sys
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


FIELDS = {
    "descripcion_item": "DESCRIPCION_ITEM",
    "descripcion_proceso": "DESCRIPCION_PROCESO",
    "entidad_nombre_original": "ENTIDAD",
    "entidad_departamento_original": "ENTIDAD_DEPARTAMENTO",
    "departamento_item_original": "DEPARTAMENTO_ITEM",
    "proveedor_nombre_original": "PROVEEDOR",
    "cantidad_decimal": "CANTIDAD_ADJUDICADO_ITEM",
    "unidad_original": "UNIDAD_MEDIDA",
    "fecha_buenapro": "FECHA_BUENAPRO",
    "estado_item": "ESTADO_ITEM",
}

REQUIRED = set(FIELDS.values()) | {
    "CODIGOCONVOCATORIA",
    "N_ITEM",
    "MONTO_ADJUDICADO_ITEM_SOLES",
}


def txt(value):
    return "" if value is None else str(value).strip()


def key(code, number):
    code = txt(code)
    if not code:
        raise ValueError("Codigo de convocatoria vacio")
    return code, int(number)


def cents(value):
    if value is None or txt(value) == "":
        raise ValueError("Monto adjudicado vacio")
    amount = Decimal(txt(value))
    if not amount.is_finite():
        raise ValueError("Monto no finito")
    return int((amount * 100).quantize(
        Decimal("1"), rounding=ROUND_HALF_UP
    ))


def decimal_text(value):
    if value is None or txt(value) == "":
        return ""
    d = Decimal(txt(value))
    if not d.is_finite():
        raise ValueError("Cantidad no finita")
    return format(d.normalize(), "f")


def fecha(value):
    if value is None or txt(value) == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(
            value, datetime
        ) else value.isoformat()
    value = txt(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Fecha no reconocida: " + value)


def read_excel(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        headers = [txt(x) for x in next(iterator)]
        missing = REQUIRED - set(headers)
        if missing:
            raise ValueError("Columnas ausentes: " + str(sorted(missing)))

        found = {}
        for line, values in enumerate(iterator, 2):
            if all(v is None for v in values):
                continue
            row = dict(zip(headers, values))
            k = key(row["CODIGOCONVOCATORIA"], row["N_ITEM"])
            if k in found:
                raise ValueError("Clave duplicada en Excel: " + str(k))
            found[k] = row
        return found
    finally:
        wb.close()


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--excel", required=True)
    p.add_argument("--json", required=True)
    p.add_argument("--aplicar", action="store_true")
    args = p.parse_args()

    path = Path(args.json)
    with path.open(encoding="utf-8-sig") as f:
        document = json.load(f)

    if not isinstance(document, dict):
        raise ValueError("Raiz JSON incorrecta")
    if not isinstance(document.get("datos"), list):
        raise ValueError("Lista datos ausente")

    records = document["datos"]
    existing = {}
    for record in records:
        k = key(record["codigo_convocatoria"], record["numero_item"])
        if k in existing:
            raise ValueError("Clave duplicada en JSON: " + str(k))
        existing[k] = record

    source = read_excel(Path(args.excel))
    old_keys = set(existing)
    new_keys = set(source)
    missing = old_keys - new_keys
    additions = new_keys - old_keys

    print("=== RADAR SIE V2.17 ===")
    print("ITEMS_JSON=", len(existing))
    print("ITEMS_EXCEL=", len(source))
    print("ITEMS_NUEVOS=", len(additions))
    print("ITEMS_AUSENTES_EN_EXCEL=", len(missing))

    if len(existing) < 6209:
        raise ValueError("JSON historico menor al minimo validado")
    if missing:
        print("HISTORICOS_CONSERVADOS=", len(missing))
        print("POLITICA=NO_ELIMINAR_HISTORICOS")

    provider_ids = {}
    max_id = 0
    for record in records:
        name = txt(record.get("proveedor_nombre_original"))
        pid = record.get("proveedor_id")
        if pid is not None:
            pid = int(pid)
            max_id = max(max_id, pid)
            if name:
                provider_ids.setdefault(name, pid)

    additions_list = []
    for k in sorted(additions):
        row = source[k]
        record = {
            "codigo_convocatoria": k[0],
            "numero_item": k[1],
        }
        for field, column in FIELDS.items():
            value = row.get(column)
            if field == "cantidad_decimal":
                value = decimal_text(value)
            elif field == "fecha_buenapro":
                value = fecha(value)
            else:
                value = txt(value)
            record[field] = value

        record["monto_adjudicado_centimos"] = cents(
            row["MONTO_ADJUDICADO_ITEM_SOLES"]
        )
        name = record["proveedor_nombre_original"]
        if name not in provider_ids:
            max_id += 1
            provider_ids[name] = max_id
        record["proveedor_id"] = provider_ids[name]
        additions_list.append(record)

    if not args.aplicar:
        print("MODO=SIMULACION")
        print("JSON_MODIFICADO=NO")
        print("RESULTADO=PASS_DRY_RUN")
        return

    if not additions_list:
        print("RESULTADO=SIN_CAMBIOS")
        return

    document["datos"].extend(additions_list)
    document["resumen"]["items"] = len(document["datos"])

    # Atomic replacement: only after all validations finish.
    temp = path.with_name(path.name + ".tmp")
    if temp.exists():
        raise ValueError("Archivo temporal existente: " + str(temp))
    try:
        with temp.open("w", encoding="utf-8", newline="\n") as f:
            json.dump(
                document, f, ensure_ascii=False,
                separators=(",", ":")
            )
            f.write("\n")
        temp.replace(path)
    finally:
        if temp.exists():
            temp.unlink()

    print("ITEMS_PUBLICABLES=", len(document["datos"]))
    print("RESULTADO=ACTUALIZACION_LOCAL_PASS")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print("RESULTADO=BLOQUEADO", file=sys.stderr)
        print("ERROR=", str(exc), file=sys.stderr)
        sys.exit(1)